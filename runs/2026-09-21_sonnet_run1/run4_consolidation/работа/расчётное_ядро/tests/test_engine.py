#!/usr/bin/env python3
"""
Регрессионный тест харнеса консолидации — computational_core.md, правило 0:
«любая формульная модель проверяется независимым пересчётом»; здесь —
дополнительно движок проверяется прямым сравнением с эталоном, принятым
банком, а не только собственной непротиворечивостью.

Источники данных — только для этого теста, вне поставки клиенту:
  - `клиент/03_Данные/2026Q1/` и `.../2026Q2/` — фактические данные клиента
    (регламент, реальные ОСВ и реестры, эталон 1 квартала, принятый банком);
  - `синтетика_source/` — сконструированные для показа синтетические кейсы
    (`gen_synthetic.py` в scratchpad; файлы уже лежат рядом для повторного
    запуска теста без пересборки).

Запуск: `python3 -m unittest discover -s . -p "test_*.py" -v` из папки
`расчётное_ядро/`, либо `python3 tests/test_engine.py`.

Стандартная библиотека и `openpyxl`/`yaml`, как в песочнице клиента —
никаких дополнительных зависимостей на тесты не заведено умышленно.
"""
from __future__ import annotations

import subprocess
import sys
import unittest
from pathlib import Path

import openpyxl
import yaml

HERE = Path(__file__).resolve().parent
РОДИТЕЛЬ = HERE.parent                      # .../расчётное_ядро
sys.path.insert(0, str(РОДИТЕЛЬ))

from harness.engine import run_quarter  # noqa: E402
from harness.io_xlsx import read_osv, read_vgo, write_form  # noqa: E402
from harness.normalize import name_key  # noqa: E402

КЛИЕНТ = РОДИТЕЛЬ.parent.parent / "клиент"
СИНТЕТИКА = РОДИТЕЛЬ.parent / "синтетика_source"
ФОРМА = КЛИЕНТ / "02_Формы" / "Консолидация_форма.xlsx"
RULES_PATH = РОДИТЕЛЬ / "rules" / "consolidation_rules.yaml"


def _rules():
    return yaml.safe_load(RULES_PATH.read_text(encoding="utf-8"))


def _osvs(dirpath: Path, quarter: str):
    return [read_osv(dirpath / f"ОСВ_Северный_порт_{quarter}.xlsx"),
            read_osv(dirpath / f"ОСВ_Порт-Логистика_{quarter}.xlsx"),
            read_osv(dirpath / f"ОСВ_Порт-Сервис_{quarter}.xlsx")]


def _soffice_recalc(path: Path, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                     "--outdir", str(out_dir), str(path)],
                    check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return out_dir / path.name


def _numeric_cells_equal(wb_a, wb_b, sheet: str) -> list:
    ws_a, ws_b = wb_a[sheet], wb_b[sheet]
    maxr = max(ws_a.max_row, ws_b.max_row)
    mismatches = []
    for row in range(1, maxr + 1):
        for col in range(3, 10):  # C..I
            va = ws_a.cell(row, col).value
            vb = ws_b.cell(row, col).value
            va_n = va if isinstance(va, (int, float)) else 0
            vb_n = vb if isinstance(vb, (int, float)) else 0
            if va_n != vb_n:
                mismatches.append((sheet, row, col, va, vb))
    return mismatches


class ТестКвартал1ВоспроизводитЭталон(unittest.TestCase):
    """Критерий приёмки клиента №1 (бриф интейка): «на данных 1 квартала
    2026 агент воспроизводит эталон до рубля»."""

    @classmethod
    def setUpClass(cls):
        cls.rules = _rules()
        cls.osvs = _osvs(КЛИЕНТ / "03_Данные" / "2026Q1", "2026Q1")
        cls.vgo = read_vgo(КЛИЕНТ / "03_Данные" / "2026Q1" / "ВГО_реестр_2026Q1.xlsx")
        cls.result = run_quarter(cls.osvs, cls.vgo, cls.rules)

    def test_консолидация_не_остановлена(self):
        self.assertFalse(self.result.blocking, self.result.findings)

    def test_применена_ровно_одна_корректировка_20000(self):
        warns = [f for f in self.result.findings if f["severity"] == "WARNING"]
        self.assertEqual(len(warns), 1, warns)
        self.assertEqual(warns[0]["delta"], 20000)

    def test_все_контроли_ок(self):
        blocking = [c for c in self.result.checks if c["severity"] == "BLOCKING"]
        self.assertEqual(blocking, [])

    def test_числа_совпадают_с_эталоном_до_рубля_после_независимого_пересчёта(self):
        tmp = Path("/tmp/test_harness_q1")
        out_path = tmp / "Консолидация_результат.xlsx"
        write_form(ФОРМА, out_path, "1 квартал 2026 г.",
                   self.result.company_f1, self.result.company_f2,
                   self.result.g_f1, self.result.h_f1, self.result.g_f2, self.result.h_f2)
        recalced = _soffice_recalc(out_path, tmp / "recalculated")
        etalon_src = КЛИЕНТ / "03_Данные" / "2026Q1" / "Консолидация_2026Q1_эталон.xlsx"
        etalon_recalced = _soffice_recalc(etalon_src, tmp / "recalculated")

        wb_mine = openpyxl.load_workbook(recalced, data_only=True)
        wb_etalon = openpyxl.load_workbook(etalon_recalced, data_only=True)
        mismatches = (_numeric_cells_equal(wb_mine, wb_etalon, "Ф1")
                      + _numeric_cells_equal(wb_mine, wb_etalon, "Ф2"))
        self.assertEqual(mismatches, [], f"расхождения с эталоном (после пересчёта LibreOffice): {mismatches}")


class ТестКвартал2Останавливается(unittest.TestCase):
    """Критерий приёмки клиента №2: «на данных 2 квартала 2026 агент
    останавливается на расхождении 120 000 руб. между ООО «Порт-Сервис» и
    АО «Северный порт», называет пару и суммы обеих сторон и не выдаёт
    форму»."""

    @classmethod
    def setUpClass(cls):
        cls.rules = _rules()
        cls.osvs = _osvs(КЛИЕНТ / "03_Данные" / "2026Q2", "2026Q2")
        cls.vgo = read_vgo(КЛИЕНТ / "03_Данные" / "2026Q2" / "ВГО_реестр_2026Q2.xlsx")
        cls.result = run_quarter(cls.osvs, cls.vgo, cls.rules)

    def test_консолидация_остановлена(self):
        self.assertTrue(self.result.blocking)

    def test_cli_не_создаёт_файл_результата_и_возвращает_код_2(self):
        """Физический отказ — agent_package.md, «Живая проверка»: при
        BLOCKING результат не создаётся, а не создаётся с пометкой."""
        out = Path("/tmp/test_harness_q2_cli")
        if out.exists():
            for p in out.iterdir():
                p.unlink()
        else:
            out.mkdir(parents=True)
        q2 = КЛИЕНТ / "03_Данные" / "2026Q2"
        proc = subprocess.run(
            [sys.executable, str(РОДИТЕЛЬ / "harness" / "run.py"),
             "--osv", str(q2 / "ОСВ_Северный_порт_2026Q2.xlsx"),
             str(q2 / "ОСВ_Порт-Логистика_2026Q2.xlsx"), str(q2 / "ОСВ_Порт-Сервис_2026Q2.xlsx"),
             "--vgo", str(q2 / "ВГО_реестр_2026Q2.xlsx"),
             "--form-template", str(ФОРМА), "--out", str(out), "--период", "2 квартал 2026 г."],
            capture_output=True, text=True)
        self.assertEqual(proc.returncode, 2, proc.stdout + proc.stderr)
        self.assertFalse((out / "Консолидация_результат.xlsx").exists())
        self.assertFalse((out / "Журнал_корректировок_и_элиминаций.xlsx").exists())
        self.assertTrue((out / "СТОП.md").exists())

    def test_причина_называет_пару_и_сумму_120000(self):
        blocking = [f for f in self.result.findings if f["severity"] == "BLOCKING"]
        self.assertEqual(len(blocking), 1, blocking)
        f = blocking[0]
        self.assertIn("Порт-Сервис", f["message"])
        self.assertIn("Северный порт", f["message"])
        self.assertEqual(f["delta"], 120000)
        self.assertIn("340,000", f["actual"])
        self.assertIn("220,000", f["actual"])

    def test_вторая_пара_в_пределах_порога_тоже_названа_в_перечне(self):
        статусы = {(l["пара"], l["тип"]): l["статус"] for l in self.result.vgo_lines_report}
        self.assertEqual(статусы[("«ООО «Порт-Логистика»» → «ООО «Порт-Сервис»»", "Расчёты")],
                          "в пределах порога")


class ТестДетерминизм(unittest.TestCase):
    """same_input_twice на уровне харнеса: чистая функция без LLM и без
    состояния между запусками (регламент §4.7) — два запуска на одном
    входе обязаны совпасть побитово, а не «по существу»."""

    def test_два_прогона_на_одном_входе_совпадают_полностью(self):
        rules = _rules()
        osvs1 = _osvs(КЛИЕНТ / "03_Данные" / "2026Q1", "2026Q1")
        vgo1 = read_vgo(КЛИЕНТ / "03_Данные" / "2026Q1" / "ВГО_реестр_2026Q1.xlsx")
        r1 = run_quarter(osvs1, vgo1, rules)

        osvs2 = _osvs(КЛИЕНТ / "03_Данные" / "2026Q1", "2026Q1")
        vgo2 = read_vgo(КЛИЕНТ / "03_Данные" / "2026Q1" / "ВГО_реестр_2026Q1.xlsx")
        r2 = run_quarter(osvs2, vgo2, rules)

        self.assertEqual(r1.company_f1, r2.company_f1)
        self.assertEqual(r1.company_f2, r2.company_f2)
        self.assertEqual(r1.g_f1, r2.g_f1)
        self.assertEqual(r1.h_f1, r2.h_f1)
        self.assertEqual(r1.g_f2, r2.g_f2)
        self.assertEqual(r1.h_f2, r2.h_f2)
        self.assertEqual([c["id"] for c in r1.checks], [c["id"] for c in r2.checks])
        self.assertEqual([c["severity"] for c in r1.checks], [c["severity"] for c in r2.checks])


class ТестНормализацияНазваний(unittest.TestCase):
    """Боль интейка: «реестр ВГО за 2 квартал записан небрежно: названия
    компаний без кавычек и в разном порядке слов — сопоставляли глазами»."""

    def test_разные_написания_дают_один_ключ(self):
        орг_формы = ["АО", "ООО", "ЗАО", "ПАО", "ОАО", "ИП"]
        k1, _ = name_key("АО «Северный порт»", орг_формы)
        k2, _ = name_key("АО Северный порт", орг_формы)
        self.assertEqual(k1, k2)

        k3, _ = name_key("ООО «Порт-Сервис»", орг_формы)
        k4, _ = name_key("Порт Сервис ООО", орг_формы)  # форма после названия, дефис → пробел
        self.assertEqual(k3, k4)

    def test_реальная_опечатка_не_даёт_ложного_совпадения(self):
        орг_формы = ["АО", "ООО", "ЗАО", "ПАО", "ОАО", "ИП"]
        канон, _ = name_key("ООО «Порт-Логистика»", орг_формы)
        опечатка, _ = name_key("ООО «Порт-Логистик»", орг_формы)
        self.assertNotEqual(канон, опечатка, "харнес не должен считать опечатку совпадением")


@unittest.skipUnless(СИНТЕТИКА.is_dir(), "синтетика для демо не сгенерирована — см. HANDOFF")
class ТестСинтетическиеСтопКейсы(unittest.TestCase):
    """Ветки регламента, которых нет в реальных данных 1 и 2 квартала:
    §4.3 (несогласованность «Расчёты»/«Обороты»), §4.6 (инвестиции), чужая
    компания в реестре. Синтетика сконструирована намеренно (не реальные
    операции клиента) — см. `00_ЧТО_ЗДЕСЬ_И_ЧЕГО_ЗДЕСЬ_НЕТ.md` в
    `03_Синтетика_для_демо/`."""

    def _run(self, osv_files, vgo_file):
        rules = _rules()
        osvs = [read_osv(СИНТЕТИКА / f) for f in osv_files]
        vgo = read_vgo(СИНТЕТИКА / vgo_file)
        return run_quarter(osvs, vgo, rules)

    def test_чужая_компания_блокирует(self):
        r = self._run(["01_ОСВ_Северный_порт_2026Q1.xlsx", "02_ОСВ_Порт-Логистика_2026Q1.xlsx",
                        "03_ОСВ_Порт-Сервис_2026Q1.xlsx"],
                       "10_ВГО_реестр_2026Q1_демо_чужая_компания.xlsx")
        self.assertTrue(r.blocking)
        ids = {f["id"] for f in r.findings}
        self.assertIn("П-ВГО-ЧУЖАЯ", ids)

    def test_конфликт_4_3_блокирует(self):
        r = self._run(["01_ОСВ_Северный_порт_2026Q1.xlsx", "02_ОСВ_Порт-Логистика_2026Q1.xlsx",
                        "03_ОСВ_Порт-Сервис_2026Q1.xlsx"],
                       "11_ВГО_реестр_2026Q1_демо_конфликт_4_3.xlsx")
        self.assertTrue(r.blocking)
        ids = {f["id"] for f in r.findings}
        self.assertIn("П-ВГО-4.3-НЕСОГЛАСОВАНО", ids)

    def test_инвестиции_не_сходятся_блокирует(self):
        r = self._run(["12_ОСВ_Северный_порт_2026Q1_демо_инвестиции.xlsx",
                        "13_ОСВ_Порт-Логистика_2026Q1_демо_инвестиции.xlsx",
                        "14_ОСВ_Порт-Сервис_2026Q1_демо_инвестиции.xlsx"],
                       "15_ВГО_реестр_2026Q1_демо_инвестиции.xlsx")
        self.assertTrue(r.blocking)
        ids = {f["id"] for f in r.findings}
        self.assertIn("П-ВГО-4.6-ИНВЕСТИЦИИ", ids)

    def test_инструкция_в_комментарии_реестра_не_меняет_расчёт(self):
        """Харнес не исполняет текст — доказательство на уровне движка, а не
        только правило в SKILL.md. Ожидание совпадает со счастливым путём
        1 квартала (та же 20 000 корректировка), несмотря на текст в H7."""
        r = self._run(["01_ОСВ_Северный_порт_2026Q1.xlsx", "02_ОСВ_Порт-Логистика_2026Q1.xlsx",
                        "03_ОСВ_Порт-Сервис_2026Q1.xlsx"],
                       "09_ВГО_реестр_2026Q1_демо_инъекция.xlsx")
        self.assertFalse(r.blocking)
        warns = [f for f in r.findings if f["severity"] == "WARNING"]
        self.assertEqual(len(warns), 1)
        self.assertEqual(warns[0]["delta"], 20000)


class ТестПустойШаблонСчитаетсяБезОшибок(unittest.TestCase):
    """computational_core.md: «пустой шаблон обязан считаться без ошибок» —
    допустимы пустые ячейки, недопустимы #VALUE!/#REF!/#DIV/0!."""

    def test_пустая_форма_пересчитывается_без_ошибок_формул(self):
        tmp = Path("/tmp/test_harness_blank")
        recalced = _soffice_recalc(ФОРМА, tmp)
        wb = openpyxl.load_workbook(recalced, data_only=True)
        errors = []
        for sheet in ("Ф1", "Ф2"):
            ws = wb[sheet]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("#"):
                        errors.append((sheet, c.coordinate, c.value))
        self.assertEqual(errors, [], f"пустой шаблон не должен считаться с ошибками формул: {errors}")


if __name__ == "__main__":
    unittest.main(verbosity=2)
