"""Чтение ОСВ и реестра ВГО, запись результата в форму консолидации.

Формат ОСВ и реестра — выгрузка 1С и утверждённая форма клиента
(`клиент/03_Данные/...`, `клиент/02_Формы/Консолидация_форма.xlsx`).
Харнес читает их как есть и не подгоняет структуру под код: если строка
«Счет» не находится или итоговая строка отсутствует, это ошибка входа
(`BLOCKING`), а не повод угадать раскладку.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl


@dataclass
class ОСВ:
    путь: str
    организация_сырая: str
    счета: dict = field(default_factory=dict)  # код счёта → {начало_дт, начало_кт, оборот_дт, оборот_кт, конец_дт, конец_кт}
    итого: dict = field(default_factory=dict)


def read_osv(path: str | Path) -> ОСВ:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    орг = None
    for row in ws.iter_rows(min_row=1, max_row=4):
        for c in row:
            if isinstance(c.value, str) and c.value.strip().startswith("Организация"):
                орг = c.value.split(":", 1)[1].strip() if ":" in c.value else c.value.strip()
    if орг is None:
        raise ValueError(f"{path}: не найдена строка «Организация:» — не ОСВ или изменён формат")

    # Строка заголовка — та, где колонка A буквально «Счет».
    header_row = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip() == "Счет":
            header_row = r
            break
    if header_row is None:
        raise ValueError(f"{path}: не найдена строка заголовка «Счет» — изменён формат ОСВ")

    data_start = header_row + 2  # заголовок + строка «Дебет/Кредит»
    счета, итого = {}, {}
    for r in range(data_start, ws.max_row + 1):
        код = ws.cell(r, 1).value
        if код is None:
            continue
        код = str(код).strip()
        vals = {
            "начало_дт": ws.cell(r, 3).value or 0,
            "начало_кт": ws.cell(r, 4).value or 0,
            "оборот_дт": ws.cell(r, 5).value or 0,
            "оборот_кт": ws.cell(r, 6).value or 0,
            "конец_дт": ws.cell(r, 7).value or 0,
            "конец_кт": ws.cell(r, 8).value or 0,
        }
        if код == "Итого":
            итого = vals
            break
        счета[код] = vals
    if not итого:
        raise ValueError(f"{path}: не найдена итоговая строка «Итого» — файл обрезан или изменён формат")
    return ОСВ(путь=str(path), организация_сырая=орг, счета=счета, итого=итого)


@dataclass
class СтрокаВГО:
    номер: int
    тип: str
    компания_а_сырая: str
    компания_б_сырая: str
    сумма_а: float
    сумма_б: float
    комментарий: str
    путь: str
    строка_листа: int


def read_vgo(path: str | Path) -> list[СтрокаВГО]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header_row = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip() == "№":
            header_row = r
            break
    if header_row is None:
        raise ValueError(f"{path}: не найдена строка заголовка реестра ВГО (столбец «№»)")

    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        num = ws.cell(r, 1).value
        if num is None:
            continue
        out.append(СтрокаВГО(
            номер=int(num),
            тип=str(ws.cell(r, 2).value or "").strip(),
            компания_а_сырая=str(ws.cell(r, 3).value or "").strip(),
            компания_б_сырая=str(ws.cell(r, 4).value or "").strip(),
            сумма_а=float(ws.cell(r, 5).value or 0),
            сумма_б=float(ws.cell(r, 6).value or 0),
            комментарий=str(ws.cell(r, 8).value or "").strip(),
            путь=str(path), строка_листа=r,
        ))
    return out


# Колонки формы Ф1/Ф2, как в утверждённом шаблоне (02_Формы/Консолидация_форма.xlsx):
# A код, B статья, C/D/E компании (SP/PL/PS в этом порядке), F сумма по группе
# (формула), G корректировки, H элиминации (обе пишет харнес), I консолидировано (формула).
КОЛОНКА_КОМПАНИИ = {"SP": 3, "PL": 4, "PS": 5}   # C, D, E
КОЛОНКА_G, КОЛОНКА_H = 7, 8


def _row_by_code(ws, code: str) -> int | None:
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
        if str(r[0].value or "").strip() == code:
            return r[0].row
    return None


def write_form(template_path: str | Path, out_path: str | Path, period_label: str,
                company_values_f1: dict, company_values_f2: dict,
                g_f1: dict, h_f1: dict, g_f2: dict, h_f2: dict) -> None:
    """Копирует утверждённую форму и вписывает значения в C/D/E (по компаниям),
    G (корректировки) и H (элиминации). Формулы F и I не трогает — они уже
    стоят в шаблоне."""
    wb = openpyxl.load_workbook(template_path)  # formulas kept (data_only=False по умолчанию)
    for sheet, company_vals, g_vals, h_vals in (
        ("Ф1", company_values_f1, g_f1, h_f1),
        ("Ф2", company_values_f2, g_f2, h_f2),
    ):
        ws = wb[sheet]
        ws["A2"] = f"Период: {period_label}"
        for code, per_company in company_vals.items():
            r = _row_by_code(ws, code)
            if r is None:
                raise ValueError(f"{template_path}: строка с кодом {code!r} не найдена на листе {sheet}")
            for comp_code, col in КОЛОНКА_КОМПАНИИ.items():
                v = per_company.get(comp_code)
                if v:
                    ws.cell(r, col, int(v))
        for code, v in g_vals.items():
            r = _row_by_code(ws, code)
            if r and v:
                ws.cell(r, КОЛОНКА_G, int(v))
        for code, v in h_vals.items():
            r = _row_by_code(ws, code)
            if r and v:
                ws.cell(r, КОЛОНКА_H, int(v))
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def write_journal(out_path: str | Path, entries: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Журнал корректировок"
    ws.append(["№", "Вид", "Компания / пара", "Содержание", "Сумма, руб."])
    for i, e in enumerate(entries, 1):
        ws.append([i, e["вид"], e["пара"], e["содержание"], int(e["сумма"])])
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def write_protocol(out_path: str | Path, checks: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Проверки"
    ws.append(["id", "Проверка", "Ожидание", "Факт", "Статус", "Δ"])
    for c in checks:
        ws.append([c["id"], c["message"], str(c.get("expected", "")), str(c.get("actual", "")),
                   c["severity"], str(c.get("delta", ""))])
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
