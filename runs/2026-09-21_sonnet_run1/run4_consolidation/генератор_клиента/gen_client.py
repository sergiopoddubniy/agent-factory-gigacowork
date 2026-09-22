#!/usr/bin/env python3
"""Синтетические материалы клиента «Группа Северный порт» для теста фабрики.

Генерирует: ОСВ трёх компаний за 1 и 2 квартал 2026, реестры ВГО, пустую
клиентскую форму консолидации, эталонную консолидацию 1 кв. (по методике).
Все числа — вымышленные. Инварианты методики соблюдаются конструктивно и
проверяются в конце скрипта (assert).
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "клиент"
THRESHOLD = 50_000

PARENT = "АО «Северный порт»"
SUB1 = "ООО «Порт-Логистика»"
SUB2 = "ООО «Порт-Сервис»"
COMPANIES = [PARENT, SUB1, SUB2]
SHORT = {PARENT: "Северный_порт", SUB1: "Порт-Логистика", SUB2: "Порт-Сервис"}

ACCOUNTS = [  # (счёт, наименование)
    ("01", "Основные средства"),
    ("02", "Амортизация основных средств"),
    ("10", "Материалы"),
    ("41", "Товары"),
    ("51", "Расчетные счета"),
    ("58", "Финансовые вложения"),
    ("60", "Расчеты с поставщиками и подрядчиками"),
    ("62", "Расчеты с покупателями и заказчиками"),
    ("67", "Расчеты по долгосрочным кредитам и займам"),
    ("68.04", "Налог на прибыль"),
    ("80", "Уставный капитал"),
    ("84", "Нераспределенная прибыль (непокрытый убыток)"),
    ("90.01", "Выручка"),
    ("90.02", "Себестоимость продаж"),
    ("90.09", "Прибыль / убыток от продаж"),
    ("91.01", "Прочие доходы"),
    ("91.02", "Прочие расходы"),
    ("91.09", "Сальдо прочих доходов и расходов"),
    ("99", "Прибыли и убытки"),
]

# --- исходные данные компаний -------------------------------------------
# Балансовые остатки на 01.01.2026 (Дт положительное, Кт отрицательное — для
# удобства генерации), обороты за квартал.
#
# Компания: dict счёт -> сальдо начало (знак: + Дт, - Кт)
OPEN_Q1 = {
    PARENT: {"01": 48_000_000, "02": -12_500_000, "10": 3_200_000, "41": 0,
             "51": 9_400_000, "58": 10_000_000, "60": -4_100_000, "62": 6_300_000,
             "67": -20_000_000, "68.04": -300_000, "80": -25_000_000, "84": -15_000_000},
    SUB1:   {"01": 21_000_000, "02": -6_200_000, "10": 1_100_000, "41": 2_400_000,
             "51": 3_150_000, "58": 0, "60": -3_800_000, "62": 4_900_000,
             "67": -12_000_000, "68.04": -150_000, "80": -6_000_000, "84": -4_400_000},
    SUB2:   {"01": 9_500_000, "02": -3_100_000, "10": 650_000, "41": 0,
             "51": 1_820_000, "58": 0, "60": -1_300_000, "62": 2_100_000,
             "67": -3_000_000, "68.04": -70_000, "80": -4_000_000, "84": -2_600_000},
}
for c, d in OPEN_Q1.items():
    assert abs(sum(d.values())) == 0, (c, sum(d.values()))

# Операции периода. cost = амортизация + списание материалов + услуги сторонних (сч. 60).
# Остатки ДЗ (62) и КЗ (60) на конец периода задаются целевыми — оплаты выводятся из них.
def period_ops(rev, cost, oth_inc, oth_exp, tax, dep, mat_used, buy, dz_close, kz_close, capex=0):
    return dict(rev=rev, cost=cost, oth_inc=oth_inc, oth_exp=oth_exp, tax=tax, dep=dep,
                mat_used=mat_used, buy=buy, dz_close=dz_close, kz_close=kz_close, capex=capex)

OPS = {
    "2026Q1": {
        PARENT: period_ops(31_200_000, 24_100_000, 410_000, 690_000, 1_364_000, 1_200_000, 7_200_000, 7_450_000, 7_400_000, 2_900_000),
        SUB1:   period_ops(17_800_000, 14_900_000, 120_000, 260_000,   552_000,   620_000, 4_500_000, 4_650_000, 5_500_000, 3_600_000),
        SUB2:   period_ops( 6_900_000,  5_300_000,  60_000, 140_000,   304_000,   310_000, 1_600_000, 1_700_000, 2_300_000, 1_200_000),
    },
    "2026Q2": {
        PARENT: period_ops(33_900_000, 26_400_000, 380_000, 720_000, 1_432_000, 1_200_000, 7_900_000, 8_000_000, 8_100_000, 3_100_000, 2_500_000),
        SUB1:   period_ops(19_400_000, 16_100_000, 150_000, 300_000,   630_000,   620_000, 4_800_000, 4_900_000, 5_900_000, 3_900_000),
        SUB2:   period_ops( 7_400_000,  5_700_000,  70_000, 160_000,   322_000,   310_000, 1_700_000, 1_750_000, 2_400_000, 1_300_000),
    },
}

# Реестр ВГО: (тип, A=продавец/кредитор, B=покупатель/дебитор, сумма по данным A, сумма по данным B, комментарий)
# «Расчёты» — остаток ДЗ у A (сч. 62) против КЗ у B (сч. 60) на конец периода.
# «Обороты» — выручка A (90.01) от B за период против себестоимости B (90.02) от A.
VGO = {
    "2026Q1": [
        ("Расчёты", PARENT, SUB1, 1_450_000, 1_450_000, "аренда причальной стенки, счёт за март"),
        ("Обороты", PARENT, SUB1, 4_350_000, 4_350_000, "аренда причальной стенки, 1 кв."),
        ("Расчёты", SUB1, SUB2, 380_000, 360_000, "перевалка, счёт от 31.03 на 20 000 не отражён у Порт-Сервис"),
        ("Обороты", SUB1, SUB2, 1_140_000, 1_120_000, "перевалка, 1 кв."),
        ("Расчёты", SUB2, PARENT, 210_000, 210_000, "обслуживание техники"),
        ("Обороты", SUB2, PARENT, 630_000, 630_000, "обслуживание техники, 1 кв."),
    ],
    "2026Q2": [
        ("Расчёты", PARENT, SUB1, 1_520_000, 1_520_000, "аренда причальной стенки, счёт за июнь"),
        ("Обороты", PARENT, SUB1, 4_560_000, 4_560_000, "аренда причальной стенки, 2 кв."),
        ("Расчёты", SUB1, SUB2, 410_000, 380_000, "перевалка; расхождение 30 000 — счёт от 30.06"),
        ("Обороты", SUB1, SUB2, 1_230_000, 1_200_000, "перевалка, 2 кв."),
        ("Расчёты", SUB2, PARENT, 340_000, 220_000, "обслуживание техники; расхождение 120 000 — акт за июнь не подписан"),
        ("Обороты", SUB2, PARENT, 720_000, 600_000, "обслуживание техники, 2 кв."),
    ],
}

# Имена компаний в реестре 2 кв. записаны небрежно (реальность клиента):
NAME_VARIANTS_Q2 = {PARENT: "АО Северный порт", SUB1: "ООО Порт-Логистика", SUB2: "Порт Сервис ООО"}


def build_osv(opening: dict, ops: dict) -> dict:
    """Возвращает по счетам: (нач Дт, нач Кт, об Дт, об Кт, кон Дт, кон Кт)."""
    dt = {a: 0 for a, _ in ACCOUNTS}
    kt = {a: 0 for a, _ in ACCOUNTS}

    def post(d, k, s):
        dt[d] += s
        kt[k] += s

    dz_open = opening.get("62", 0)
    kz_open = -opening.get("60", 0)
    services = ops["cost"] - ops["dep"] - ops["mat_used"]
    assert services > 0, ("услуги", services)
    # начисление выручки: Дт 62 Кт 90.01
    post("62", "90.01", ops["rev"])
    # себестоимость: амортизация Дт 90.02 Кт 02; материалы Дт 90.02 Кт 10; услуги Дт 90.02 Кт 60
    post("90.02", "02", ops["dep"])
    post("90.02", "10", ops["mat_used"])
    post("90.02", "60", services)
    # закупка материалов: Дт 10 Кт 60
    post("10", "60", ops["buy"])
    # капвложения: Дт 01 Кт 60
    if ops.get("capex"):
        post("01", "60", ops["capex"])
    # прочие доходы: Дт 51 Кт 91.01 ; прочие расходы: Дт 91.02 Кт 51
    post("51", "91.01", ops["oth_inc"])
    post("91.02", "51", ops["oth_exp"])
    # оплаты: из целевых остатков ДЗ и КЗ
    paid_in = dz_open + ops["rev"] - ops["dz_close"]
    paid_out = kz_open + services + ops["buy"] + ops.get("capex", 0) - ops["kz_close"]
    assert paid_in > 0 and paid_out > 0, (paid_in, paid_out)
    post("51", "62", paid_in)
    post("60", "51", paid_out)
    # закрытие 90: Дт 90.09 Кт 99 на (rev - cost)
    post("90.09", "99", ops["rev"] - ops["cost"])
    # закрытие 91
    net91 = ops["oth_inc"] - ops["oth_exp"]
    if net91 >= 0:
        post("91.09", "99", net91)
    else:
        post("99", "91.09", -net91)
    # налог на прибыль: Дт 99 Кт 68.04 ; уплата налога прошлого периода: Дт 68.04 Кт 51
    post("99", "68.04", ops["tax"])
    open_tax = -opening.get("68.04", 0)
    post("68.04", "51", open_tax)

    res = {}
    for a, _ in ACCOUNTS:
        o = opening.get(a, 0)
        odt, okt = (o, 0) if o >= 0 else (0, -o)
        close = o + dt[a] - kt[a]
        cdt, ckt = (close, 0) if close >= 0 else (0, -close)
        res[a] = (odt, okt, dt[a], kt[a], cdt, ckt)
    s = sum(v[4] - v[5] for v in res.values())
    assert s == 0, ("ОСВ не сходится", s)
    return res


def write_osv(path: Path, company: str, period_title: str, osv: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "ОСВ"
    bold = Font(bold=True)
    ws["A1"] = f"Оборотно-сальдовая ведомость за {period_title}"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = f"Организация: {company}"
    ws["A3"] = "Выводимые данные: БУ (данные бухгалтерского учета). Единица измерения: руб."
    hdr = ["Счет", "Наименование счета", "Сальдо на начало периода", "", "Обороты за период", "", "Сальдо на конец периода", ""]
    sub = ["", "", "Дебет", "Кредит", "Дебет", "Кредит", "Дебет", "Кредит"]
    ws.append([])
    ws.append(hdr)
    ws.append(sub)
    for c in range(1, 9):
        ws.cell(row=5, column=c).font = bold
        ws.cell(row=6, column=c).font = bold
        ws.cell(row=5, column=c).alignment = Alignment(horizontal="center")
        ws.cell(row=6, column=c).alignment = Alignment(horizontal="center")
    ws.merge_cells("C5:D5"); ws.merge_cells("E5:F5"); ws.merge_cells("G5:H5")
    tot = [0] * 6
    for a, name in ACCOUNTS:
        v = osv[a]
        row = [a, name] + [x if x else None for x in v]
        ws.append(row)
        for i in range(6):
            tot[i] += v[i]
    ws.append(["Итого", ""] + tot)
    for c in range(1, 9):
        ws.cell(row=ws.max_row, column=c).font = bold
    for r in range(7, ws.max_row + 1):
        for c in range(3, 9):
            ws.cell(row=r, column=c).number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 44
    for c in "CDEFGH":
        ws.column_dimensions[c].width = 18
    wb.save(path)


def write_vgo(path: Path, period_title: str, rows, names=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр ВГО"
    ws["A1"] = f"Реестр внутригрупповых операций за {period_title}"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Составил: отдел консолидации. Суммы в руб. Расчёты — остатки на конец периода (сч. 62 у А / сч. 60 у Б); Обороты — за период (сч. 90.01 у А / сч. 90.02 у Б)."
    ws.append([])
    ws.append(["№", "Тип", "Компания А (продавец / кредитор)", "Компания Б (покупатель / дебитор)",
               "Сумма по данным А", "Сумма по данным Б", "Расхождение (А − Б)", "Комментарий"])
    for c in range(1, 9):
        ws.cell(row=4, column=c).font = Font(bold=True)
    for i, (t, a, b, sa, sb, com) in enumerate(rows, 1):
        na, nb = (names or {}).get(a, a), (names or {}).get(b, b)
        ws.append([i, t, na, nb, sa, sb, f"=E{4 + i}-F{4 + i}", com])
    for r in range(5, ws.max_row + 1):
        for c in (5, 6, 7):
            ws.cell(row=r, column=c).number_format = "#,##0.00"
    widths = [4, 10, 30, 30, 18, 18, 18, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)


F1_ROWS = [  # (код, статья, тип): тип 'v' — вводится, 's' — итог (формула суммы кодов), 'z' — статья без источника (переносится как есть)
    ("1150", "Основные средства", "v"),
    ("1170", "Финансовые вложения", "v"),
    ("1190", "Прочие внеоборотные активы", "z"),
    ("1100", "Итого по разделу I", "s:1150,1170,1190"),
    ("1210", "Запасы", "v"),
    ("1230", "Дебиторская задолженность", "v"),
    ("1250", "Денежные средства и денежные эквиваленты", "v"),
    ("1260", "Прочие оборотные активы", "z"),
    ("1200", "Итого по разделу II", "s:1210,1230,1250,1260"),
    ("1600", "БАЛАНС (актив)", "s:1100,1200"),
    ("1310", "Уставный капитал", "v"),
    ("1370", "Нераспределенная прибыль (непокрытый убыток)", "v"),
    ("1300", "Итого по разделу III", "s:1310,1370"),
    ("1410", "Заемные средства (долгосрочные)", "v"),
    ("1400", "Итого по разделу IV", "s:1410"),
    ("1520", "Кредиторская задолженность", "v"),
    ("1540", "Оценочные обязательства", "z"),
    ("1500", "Итого по разделу V", "s:1520,1540"),
    ("1700", "БАЛАНС (пассив)", "s:1300,1400,1500"),
]
F2_ROWS = [
    ("2110", "Выручка", "v"),
    ("2120", "Себестоимость продаж", "v"),
    ("2200", "Прибыль (убыток) от продаж", "d:2110-2120"),
    ("2340", "Прочие доходы", "v"),
    ("2350", "Прочие расходы", "v"),
    ("2300", "Прибыль (убыток) до налогообложения", "d:2200+2340-2350"),
    ("2410", "Налог на прибыль", "v"),
    ("2400", "Чистая прибыль (убыток)", "d:2300-2410"),
]
COLS = ["Код", "Статья", PARENT, SUB1, SUB2, "Сумма по группе", "Корректировки", "Элиминации", "Консолидировано"]


def write_form(path: Path, period_title: str, values=None, adjustments=None):
    """Клиентская форма консолидации. values: {код: {компания: число}}; adjustments: {код: (корр, элим)}"""
    wb = Workbook()
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    tot_fill = PatternFill("solid", fgColor="F2F2F2")
    for sheet, rows, title in (("Ф1", F1_ROWS, "Бухгалтерский баланс (консолидированный), руб."),
                               ("Ф2", F2_ROWS, "Отчет о финансовых результатах (консолидированный), руб.")):
        ws = wb.create_sheet(sheet)
        ws["A1"] = f"Группа «Северный порт» — {title}"
        ws["A1"].font = Font(bold=True, size=12)
        ws["A2"] = f"Период: {period_title}"
        ws.append([])
        ws.append(COLS)
        for c in range(1, len(COLS) + 1):
            cell = ws.cell(row=4, column=c)
            cell.font = Font(bold=True); cell.fill = head_fill; cell.border = border
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        code_row = {}
        for i, (code, name, kind) in enumerate(rows):
            r = 5 + i
            code_row[code] = r
            ws.cell(row=r, column=1, value=code)
            ws.cell(row=r, column=2, value=name)
            if kind == "v" or kind == "z":
                for j, comp in enumerate(COMPANIES):
                    v = (values or {}).get(code, {}).get(comp)
                    ws.cell(row=r, column=3 + j, value=v)
                adj = (adjustments or {}).get(code, (None, None))
                ws.cell(row=r, column=7, value=adj[0])
                ws.cell(row=r, column=8, value=adj[1])
            else:
                for col in range(3, 10):
                    ws.cell(row=r, column=col).fill = tot_fill
            ws.cell(row=r, column=6, value=f"=SUM(C{r}:E{r})")
            ws.cell(row=r, column=9, value=f"=F{r}+G{r}-H{r}")
            if kind.startswith("s:") or kind.startswith("d:"):
                for col in range(3, 10):
                    L = get_column_letter(col)
                    if kind.startswith("s:"):
                        parts = kind[2:].split(",")
                        ws.cell(row=r, column=col, value="=" + "+".join(f"{L}{code_row[p]}" for p in parts))
                    else:
                        expr = kind[2:]
                        f = "="
                        import re
                        for tok in re.findall(r"[+-]?\d{4}", expr):
                            sign = tok[0] if tok[0] in "+-" else "+"
                            c4 = tok.lstrip("+-")
                            f += f"{sign}{L}{code_row[c4]}"
                        ws.cell(row=r, column=col, value=f.replace("=+", "="))
                    ws.cell(row=r, column=col).font = Font(bold=True)
                ws.cell(row=r, column=1).font = Font(bold=True)
                ws.cell(row=r, column=2).font = Font(bold=True)
            for col in range(1, 10):
                ws.cell(row=r, column=col).border = border
                if col >= 3:
                    ws.cell(row=r, column=col).number_format = "#,##0.00"
        last = 5 + len(rows) - 1
        if sheet == "Ф1":
            ws.cell(row=last + 2, column=2, value="Контроль: актив − пассив (должно быть 0)")
            for col in range(3, 10):
                L = get_column_letter(col)
                ws.cell(row=last + 2, column=col, value=f"={L}{code_row['1600']}-{L}{code_row['1700']}")
                ws.cell(row=last + 2, column=col).number_format = "#,##0.00"
        else:
            ws.cell(row=last + 2, column=2, value="Контроль: чистая прибыль по Ф2 − прирост нераспределённой прибыли (справочно, заполняется вручную)")
        ws.column_dimensions["A"].width = 7
        ws.column_dimensions["B"].width = 46
        for c in "CDEFGHI":
            ws.column_dimensions[c].width = 19
        ws.freeze_panes = "C5"
    del wb["Sheet"]
    return wb


def form_values(osvs: dict) -> dict:
    """Статьи форм по компаниям из ОСВ (методика, раздел 3)."""
    vals = {}
    def put(code, comp, v):
        vals.setdefault(code, {})[comp] = v
    for comp, o in osvs.items():
        cl = lambda a: o[a][4] - o[a][5]  # конечный остаток (+Дт, −Кт)
        put("1150", comp, cl("01") + cl("02"))
        put("1170", comp, cl("58"))
        put("1190", comp, 0)
        put("1210", comp, cl("10") + cl("41"))
        put("1230", comp, cl("62"))
        put("1250", comp, cl("51"))
        put("1260", comp, 0)
        put("1310", comp, -cl("80"))
        put("1370", comp, -cl("84") - cl("99"))
        put("1410", comp, -cl("67"))
        put("1520", comp, -cl("60") - cl("68.04"))
        put("1540", comp, 0)
        put("2110", comp, o["90.01"][3])       # оборот Кт 90.01
        put("2120", comp, o["90.02"][2])       # оборот Дт 90.02
        put("2340", comp, o["91.01"][3])
        put("2350", comp, o["91.02"][2])
        put("2410", comp, o["68.04"][3])       # оборот Кт 68.04 (начислен налог)
    return vals


def consolidate(vals: dict, vgo: list):
    """Методика: корректировки стороны Б в пределах порога, элиминации по данным А, инвестиции."""
    adj = {}   # код -> корректировка (со знаком, к сумме по группе)
    elim = {}  # код -> элиминация (вычитается)
    journal = []
    blocking = []
    pairs = {}
    for t, a, b, sa, sb, com in vgo:
        pairs.setdefault((a, b), {})[t] = (sa, sb, com)
    for (a, b), d in pairs.items():
        dr = d["Расчёты"][0] - d["Расчёты"][1]
        do = d["Обороты"][0] - d["Обороты"][1]
        if abs(dr) > THRESHOLD or abs(do) > THRESHOLD:
            blocking.append((a, b, dr, do))
            continue
        if dr != do:
            blocking.append((a, b, dr, do))
            continue
        if dr:
            # сторона Б доначисляет: себестоимость +δ, кредиторка +δ, прибыль −δ
            adj["2120"] = adj.get("2120", 0) + dr
            adj["1520"] = adj.get("1520", 0) + dr
            adj["1370"] = adj.get("1370", 0) - dr
            journal.append(("Корректировка", b, f"Доначисление по операциям с {a}: себестоимость +{dr:,}; кредиторская задолженность +{dr:,}; нераспределённая прибыль −{dr:,}".replace(",", " "), dr))
        # элиминации по данным А
        ra, oa = d["Расчёты"][0], d["Обороты"][0]
        elim["1230"] = elim.get("1230", 0) + ra
        elim["1520"] = elim.get("1520", 0) + ra
        elim["2110"] = elim.get("2110", 0) + oa
        elim["2120"] = elim.get("2120", 0) + oa
        journal.append(("Элиминация", f"{a} → {b}", f"ДЗ/КЗ {ra:,}; выручка/себестоимость {oa:,}".replace(",", " "), ra))
    # инвестиции материнской = УК дочерних
    inv = vals["1170"][PARENT]
    uk = vals["1310"][SUB1] + vals["1310"][SUB2]
    if inv != uk:
        blocking.append(("инвестиции", "УК", inv, uk))
    else:
        elim["1170"] = inv
        elim["1310"] = uk
        journal.append(("Элиминация", "инвестиции ↔ уставный капитал", f"1170 у материнской −{inv:,}; 1310 у дочерних −{uk:,}".replace(",", " "), inv))
    return adj, elim, journal, blocking


def write_reference(path: Path, period_title: str, vals, adj, elim, journal, checks):
    adjustments = {}
    for code in set(list(adj) + list(elim)):
        adjustments[code] = (adj.get(code), elim.get(code))
    wb = write_form(path, period_title, vals, adjustments)
    ws = wb.create_sheet("Журнал корректировок")
    ws.append(["№", "Вид", "Компания / пара", "Содержание", "Сумма, руб."])
    for c in range(1, 6):
        ws.cell(row=1, column=c).font = Font(bold=True)
    for i, (kind, who, what, s) in enumerate(journal, 1):
        ws.append([i, kind, who, what, s])
    ws.column_dimensions["B"].width = 14; ws.column_dimensions["C"].width = 40; ws.column_dimensions["D"].width = 110; ws.column_dimensions["E"].width = 16
    ws2 = wb.create_sheet("Проверки")
    ws2.append(["Проверка", "Ожидание", "Факт", "Статус"])
    for c in range(1, 5):
        ws2.cell(row=1, column=c).font = Font(bold=True)
    for name, exp, act in checks:
        ws2.append([name, exp, act, "OK" if exp == act else "СТОП"])
    ws2.column_dimensions["A"].width = 70; ws2.column_dimensions["B"].width = 18; ws2.column_dimensions["C"].width = 18
    ws3 = wb.create_sheet("Итоги")
    ws3.append(["Код", "Статья", "Консолидировано, руб."])
    for c in range(1, 4):
        ws3.cell(row=1, column=c).font = Font(bold=True)
    for code, name, v in checks_totals(vals, adj, elim):
        ws3.append([code, name, v])
    ws3.column_dimensions["B"].width = 46; ws3.column_dimensions["C"].width = 22
    wb.save(path)


def consolidated_values(vals, adj, elim):
    cons = {}
    for code in vals:
        s = sum(vals[code].values())
        cons[code] = s + adj.get(code, 0) - elim.get(code, 0)
    cons["1100"] = cons["1150"] + cons["1170"] + cons["1190"]
    cons["1200"] = cons["1210"] + cons["1230"] + cons["1250"] + cons["1260"]
    cons["1600"] = cons["1100"] + cons["1200"]
    cons["1300"] = cons["1310"] + cons["1370"]
    cons["1400"] = cons["1410"]
    cons["1500"] = cons["1520"] + cons["1540"]
    cons["1700"] = cons["1300"] + cons["1400"] + cons["1500"]
    cons["2200"] = cons["2110"] - cons["2120"]
    cons["2300"] = cons["2200"] + cons["2340"] - cons["2350"]
    cons["2400"] = cons["2300"] - cons["2410"]
    return cons


def checks_totals(vals, adj, elim):
    cons = consolidated_values(vals, adj, elim)
    names = {c: n for c, n, _ in F1_ROWS + F2_ROWS}
    order = [c for c, _, _ in F1_ROWS] + [c for c, _, _ in F2_ROWS]
    return [(c, names[c], cons[c]) for c in order]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "02_Формы").mkdir(exist_ok=True)
    period_titles = {"2026Q1": "1 квартал 2026 г.", "2026Q2": "2 квартал 2026 г."}
    opening = {c: dict(v) for c, v in OPEN_Q1.items()}
    for period in ("2026Q1", "2026Q2"):
        d = OUT / "03_Данные" / period
        d.mkdir(parents=True, exist_ok=True)
        osvs = {}
        for comp in COMPANIES:
            osv = build_osv(opening[comp], OPS[period][comp])
            osvs[comp] = osv
            write_osv(d / f"ОСВ_{SHORT[comp]}_{period}.xlsx", comp, period_titles[period], osv)
            # остатки на следующий квартал
            opening[comp] = {a: v[4] - v[5] for a, v in osv.items()}
        names = NAME_VARIANTS_Q2 if period == "2026Q2" else None
        write_vgo(d / f"ВГО_реестр_{period}.xlsx", period_titles[period], VGO[period], names)
        vals = form_values(osvs)
        # инварианты по компаниям
        for comp in COMPANIES:
            a = vals["1150"][comp] + vals["1170"][comp] + vals["1210"][comp] + vals["1230"][comp] + vals["1250"][comp]
            p = vals["1310"][comp] + vals["1370"][comp] + vals["1410"][comp] + vals["1520"][comp]
            assert a == p, (comp, a, p)
            net = vals["2110"][comp] - vals["2120"][comp] + vals["2340"][comp] - vals["2350"][comp] - vals["2410"][comp]
            o = osvs[comp]
            assert net == o["99"][3] - o["99"][2], (comp, net, o["99"])
            # ВГО не больше статьи
        adj, elim, journal, blocking = consolidate(vals, VGO[period])
        if period == "2026Q1":
            assert not blocking, blocking
            cons = consolidated_values(vals, adj, elim)
            checks = [
                ("Актив = Пассив (консолидировано)", cons["1600"], cons["1700"]),
                ("1230 конс. = Σ1230 − ВГО ДЗ", cons["1230"], sum(vals["1230"].values()) - elim["1230"]),
                ("1520 конс. = Σ1520 + корректировки − ВГО КЗ", cons["1520"], sum(vals["1520"].values()) + adj.get("1520", 0) - elim["1520"]),
                ("2110 конс. = Σ2110 − ВГО выручка", cons["2110"], sum(vals["2110"].values()) - elim["2110"]),
                ("2400 конс. = Σ2400 − Σ корректировок", cons["2400"],
                 sum(vals["2110"][c] - vals["2120"][c] + vals["2340"][c] - vals["2350"][c] - vals["2410"][c] for c in COMPANIES) - adj.get("2120", 0)),
                ("1170 материнской = Σ1310 дочерних", vals["1170"][PARENT], vals["1310"][SUB1] + vals["1310"][SUB2]),
                ("1370 конс. − Σ1370 = −Σ корректировок", cons["1370"] - sum(vals["1370"].values()), adj.get("1370", 0)),
            ]
            assert all(e == a for _, e, a in checks), checks
            write_reference(d / f"Консолидация_{period}_эталон.xlsx", period_titles[period], vals, adj, elim, journal, checks)
            print("Q1 консолидировано:", {k: v for k, v in cons.items() if k in ("1600", "1700", "2110", "2400", "1370")})
            print("Q1 журнал:", *journal, sep="\n  ")
        else:
            print("Q2 blocking:", blocking)
            cons = consolidated_values(vals, adj, elim)
    # пустая форма
    wb = write_form(OUT / "02_Формы" / "Консолидация_форма.xlsx", "<период>")
    wb.save(OUT / "02_Формы" / "Консолидация_форма.xlsx")
    print("готово:", OUT)


if __name__ == "__main__":
    main()
